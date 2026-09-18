#!/usr/bin/python3
# -*- coding: utf-8 -*-
"""Audit current sales forecast accuracy at SPU-month granularity.

Two forecast views are compared for the past N completed months:
1) persisted forecast: rows already retained in `预测对比表`;
2) leakage-free replay: rerun the *current* v4 algorithm as of the first day of
   each target month, using only sales data from months that had fully completed
   before that target month.

Actual sales are read from `销量统计_msku月度`, the same source used by the
production forecast pipeline, so the forecast-vs-actual comparison is on a
consistent quantity definition.
"""
from __future__ import annotations

import argparse
import calendar
import csv
import math
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from common.database import db_cursor
from jobs.feishu.forecast_sales_improved import compute_forecast_for_shop, load_spu_season_map
from jobs.feishu.generate_forecast_comparison import EXCLUDED_SHOPS, extract_spu_from_sku

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except Exception:
    Workbook = None
    Font = None


def month_start(d: date) -> date:
    return d.replace(day=1)


def add_months(d: date, delta: int) -> date:
    y = d.year + (d.month - 1 + delta) // 12
    m = (d.month - 1 + delta) % 12 + 1
    return date(y, m, 1)


def default_target_months(n: int) -> List[date]:
    cur = month_start(datetime.now().date())
    return [add_months(cur, -i) for i in range(n, 0, -1)]


def parse_months(text: str) -> List[date]:
    result: List[date] = []
    for part in text.split(','):
        part = part.strip()
        if not part:
            continue
        result.append(datetime.strptime(part, '%Y-%m').date().replace(day=1))
    if not result:
        raise ValueError('months is empty')
    return sorted(set(result))


def forecast_label(d: date) -> str:
    return f"{str(d.year)[-2:]}年{d.month}月预计销量"


def actual_sales_label(d: date) -> str:
    return f"{str(d.year)[-2:]}年{d.month}月销量"


def _parse_stat_month(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    try:
        s = str(value)[:10]
        dt = datetime.strptime(s, '%Y-%m-%d')
        return date(dt.year, dt.month, 1)
    except Exception:
        return None


def table_exists(name: str) -> bool:
    with db_cursor() as cur:
        cur.execute(
            """SELECT COUNT(*) AS cnt FROM information_schema.TABLES
               WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s""",
            (name,),
        )
        row = cur.fetchone() or {}
    return int(row.get('cnt', 0) or 0) > 0


def column_exists(table: str, column: str) -> bool:
    with db_cursor() as cur:
        cur.execute(
            """SELECT COUNT(*) AS cnt FROM information_schema.COLUMNS
               WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s AND COLUMN_NAME=%s""",
            (table, column),
        )
        row = cur.fetchone() or {}
    return int(row.get('cnt', 0) or 0) > 0


def read_actual_sku_months(start: date, end: date) -> List[Dict[str, Any]]:
    table = '销量统计_msku月度'
    if not table_exists(table):
        raise RuntimeError(f'缺少表: {table}')
    has_spu = column_exists(table, 'SPU')
    spu_select = ', SPU' if has_spu else ''
    shop_ph = ','.join(['%s'] * len(EXCLUDED_SHOPS))
    sql = f"""
        SELECT SKU, 店铺, 统计日期, 销量 {spu_select}
        FROM `{table}`
        WHERE 统计日期 >= %s AND 统计日期 < %s
          AND 店铺 IS NOT NULL AND 店铺 NOT IN ('', '无')
          AND SKU IS NOT NULL AND SKU NOT IN ('', '无')
          AND 店铺 NOT IN ({shop_ph})
    """
    params: List[Any] = [start.isoformat(), end.isoformat(), *sorted(EXCLUDED_SHOPS)]
    with db_cursor() as cur:
        cur.execute(sql, params)
        return list(cur.fetchall() or [])


def build_actual_maps(rows: Sequence[Dict[str, Any]]) -> Tuple[
    Dict[Tuple[str, str, date], int],
    Dict[Tuple[str, date], int],
]:
    by_shop: Dict[Tuple[str, str, date], int] = defaultdict(int)
    by_spu: Dict[Tuple[str, date], int] = defaultdict(int)
    for r in rows:
        sku = str(r.get('SKU') or '').strip()
        shop = str(r.get('店铺') or '').strip()
        d = _parse_stat_month(r.get('统计日期'))
        source_spu = str(r.get('SPU') or '').strip()
        spu = source_spu or extract_spu_from_sku(sku)
        if not sku or not shop or not spu or not d:
            continue
        qty = int(r.get('销量') or 0)
        by_shop[(spu, shop, d)] += qty
        by_spu[(spu, d)] += qty
    return dict(by_shop), dict(by_spu)


def read_persisted_forecast(target_months: Sequence[date]) -> Dict[Tuple[str, str, date], int]:
    table = '预测对比表'
    if not table_exists(table):
        return {}
    start = min(target_months)
    end = add_months(max(target_months), 1)
    shop_ph = ','.join(['%s'] * len(EXCLUDED_SHOPS))
    sql = f"""
        SELECT SPU, 店铺, 统计日期, 系统预测销量
        FROM `{table}`
        WHERE 统计日期 >= %s AND 统计日期 < %s
          AND SPU IS NOT NULL AND SPU != ''
          AND 店铺 IS NOT NULL AND 店铺 != ''
          AND 店铺 NOT IN ({shop_ph})
    """
    params: List[Any] = [start.isoformat(), end.isoformat(), *sorted(EXCLUDED_SHOPS)]
    result: Dict[Tuple[str, str, date], int] = defaultdict(int)
    with db_cursor() as cur:
        cur.execute(sql, params)
        for r in cur.fetchall() or []:
            d = _parse_stat_month(r.get('统计日期'))
            if not d:
                continue
            key = (str(r.get('SPU') or '').strip(), str(r.get('店铺') or '').strip(), d)
            result[key] += int(r.get('系统预测销量') or 0)
    return dict(result)


def _read_history_before(target: date) -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Build input for v4 using only fully completed months before target."""
    start = date(target.year - 1, 1, 1)
    rows = read_actual_sku_months(start, target)
    result: Dict[str, Dict[str, Dict[str, Any]]] = defaultdict(lambda: defaultdict(dict))
    for r in rows:
        shop = str(r.get('店铺') or '').strip()
        sku = str(r.get('SKU') or '').strip()
        d = _parse_stat_month(r.get('统计日期'))
        if not shop or not sku or not d or d >= target:
            continue
        sku_data = result[shop][sku]
        source_spu = str(r.get('SPU') or '').strip()
        if 'SPU' not in sku_data:
            sku_data['SPU'] = source_spu or extract_spu_from_sku(sku)
        label = actual_sales_label(d)
        sku_data[label] = int(sku_data.get(label, 0) or 0) + int(r.get('销量') or 0)
    return {shop: dict(skus) for shop, skus in result.items()}


def replay_month(target: date, season_map: Dict[str, str]) -> Tuple[
    Dict[Tuple[str, str, date], int],
    Dict[str, int],
]:
    shop_history = _read_history_before(target)
    result: Dict[Tuple[str, str, date], int] = defaultdict(int)
    method_counts: Dict[str, int] = defaultdict(int)
    label = forecast_label(target)
    current_date = datetime(target.year, target.month, 1)

    for shop, shop_data in shop_history.items():
        forecasts = compute_forecast_for_shop(
            shop_data,
            [label],
            current_date=current_date,
            spu_season_map=season_map,
        )
        for sku, values in forecasts.items():
            spu = str(shop_data[sku].get('SPU') or extract_spu_from_sku(sku)).strip()
            if not spu:
                continue
            qty = int(values.get(label, 0) or 0)
            result[(spu, shop, target)] += qty
            method = str(values.get('预测方法') or '')
            if ':' in method:
                method = method.split(':', 1)[1]
            family = method.split('(', 1)[0].split('[', 1)[0] or 'UNKNOWN'
            method_counts[family] += 1
    return dict(result), dict(method_counts)


def aggregate_spu_month(shop_map: Dict[Tuple[str, str, date], int]) -> Dict[Tuple[str, date], int]:
    out: Dict[Tuple[str, date], int] = defaultdict(int)
    for (spu, _shop, d), qty in shop_map.items():
        out[(spu, d)] += int(qty or 0)
    return dict(out)


def smape(forecast: float, actual: float) -> float:
    denom = abs(forecast) + abs(actual)
    if denom == 0:
        return 0.0
    return 2 * abs(forecast - actual) / denom


def build_detail(
    actual: Dict[Tuple[str, date], int],
    persisted: Dict[Tuple[str, date], int],
    replay: Dict[Tuple[str, date], int],
    months: Sequence[date],
) -> List[Dict[str, Any]]:
    keys = set(actual) | set(persisted) | set(replay)
    target_set = set(months)
    rows: List[Dict[str, Any]] = []
    for spu, d in sorted(keys, key=lambda x: (x[1], x[0])):
        if d not in target_set:
            continue
        a = int(actual.get((spu, d), 0) or 0)
        p = int(persisted.get((spu, d), 0) or 0)
        r = int(replay.get((spu, d), 0) or 0)
        rows.append({
            '月份': d.strftime('%Y-%m'),
            'SPU': spu,
            '实际销量': a,
            '历史留存预测': p,
            '历史留存误差': p - a,
            '历史留存绝对误差': abs(p - a),
            '历史留存APE': None if a == 0 else round(abs(p - a) / a, 6),
            '当前v4月初回放预测': r,
            'v4回放误差': r - a,
            'v4回放绝对误差': abs(r - a),
            'v4回放APE': None if a == 0 else round(abs(r - a) / a, 6),
        })
    return rows


def calc_metrics(rows: Sequence[Dict[str, Any]], forecast_col: str) -> Dict[str, Any]:
    if not rows:
        return {}
    actuals = [float(r['实际销量'] or 0) for r in rows]
    forecasts = [float(r[forecast_col] or 0) for r in rows]
    errors = [f - a for f, a in zip(forecasts, actuals)]
    abs_errors = [abs(e) for e in errors]
    ape = [abs(f - a) / a for f, a in zip(forecasts, actuals) if a > 0]
    smapes = [smape(f, a) for f, a in zip(forecasts, actuals)]
    total_actual = sum(actuals)
    total_forecast = sum(forecasts)
    nonzero = sum(1 for a in actuals if a > 0)
    within20 = sum(1 for f, a in zip(forecasts, actuals) if a > 0 and abs(f-a)/a <= 0.20)
    within30 = sum(1 for f, a in zip(forecasts, actuals) if a > 0 and abs(f-a)/a <= 0.30)
    within50 = sum(1 for f, a in zip(forecasts, actuals) if a > 0 and abs(f-a)/a <= 0.50)
    return {
        '记录数': len(rows),
        '实际销量合计': int(total_actual),
        '预测销量合计': int(total_forecast),
        '总量偏差': int(total_forecast - total_actual),
        'Bias%': None if total_actual == 0 else (total_forecast - total_actual) / total_actual,
        'WAPE': None if total_actual == 0 else sum(abs_errors) / total_actual,
        'MAPE(实际>0)': None if not ape else sum(ape) / len(ape),
        'sMAPE': sum(smapes) / len(smapes),
        'MAE': sum(abs_errors) / len(abs_errors),
        'RMSE': math.sqrt(sum(e * e for e in errors) / len(errors)),
        '±20%命中率': None if nonzero == 0 else within20 / nonzero,
        '±30%命中率': None if nonzero == 0 else within30 / nonzero,
        '±50%命中率': None if nonzero == 0 else within50 / nonzero,
    }


def grouped_metrics(rows: Sequence[Dict[str, Any]], forecast_col: str, key_col: str) -> List[Dict[str, Any]]:
    groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for r in rows:
        groups[str(r[key_col])].append(r)
    out = []
    for key in sorted(groups):
        m = calc_metrics(groups[key], forecast_col)
        out.append({key_col: key, **m})
    return out


def sales_bucket(actual: int) -> str:
    if actual <= 0:
        return '0'
    if actual < 10:
        return '1-9'
    if actual < 50:
        return '10-49'
    if actual < 200:
        return '50-199'
    return '200+'


def add_buckets(rows: List[Dict[str, Any]]) -> None:
    for r in rows:
        r['实际销量档位'] = sales_bucket(int(r['实际销量'] or 0))


def normalise_metric_rows(rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out = []
    for r in rows:
        out.append({k: (round(v, 6) if isinstance(v, float) else v) for k, v in r.items()})
    return out


def write_csv(path: Path, rows: Sequence[Dict[str, Any]]) -> None:
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0].keys())
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path: Path, sheets: Sequence[Tuple[str, Sequence[Dict[str, Any]]]]) -> None:
    if Workbook is None:
        return
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for title, rows in sheets:
        ws = wb.create_sheet(title[:31])
        if not rows:
            ws.append(['无数据'])
            continue
        fields = list(rows[0].keys())
        ws.append(fields)
        if Font:
            for c in ws[1]:
                c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(k) for k in fields])
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            width = min(max(len(str(c.value or '')) for c in col) + 2, 45)
            ws.column_dimensions[col[0].column_letter].width = max(width, 10)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description='SPU-month sales forecast backtest audit')
    ap.add_argument('--months', help='comma-separated YYYY-MM, default=last 3 completed months')
    ap.add_argument('--n', type=int, default=3, help='number of completed months when --months omitted')
    ap.add_argument('--output-dir', default='reports_analysis/forecast_audit')
    args = ap.parse_args()

    months = parse_months(args.months) if args.months else default_target_months(args.n)
    print('目标月份:', ', '.join(d.strftime('%Y-%m') for d in months))

    actual_rows = read_actual_sku_months(min(months), add_months(max(months), 1))
    actual_shop, actual_spu = build_actual_maps(actual_rows)

    persisted_shop = read_persisted_forecast(months)
    persisted_spu = aggregate_spu_month(persisted_shop)

    season_map = load_spu_season_map()
    replay_shop: Dict[Tuple[str, str, date], int] = {}
    method_rows: List[Dict[str, Any]] = []
    for target in months:
        month_map, method_counts = replay_month(target, season_map)
        replay_shop.update(month_map)
        for method, count in sorted(method_counts.items(), key=lambda x: (-x[1], x[0])):
            method_rows.append({'月份': target.strftime('%Y-%m'), '预测方法族': method, 'SKU数': count})
    replay_spu = aggregate_spu_month(replay_shop)

    detail = build_detail(actual_spu, persisted_spu, replay_spu, months)
    add_buckets(detail)

    overall = normalise_metric_rows([
        {'口径': '历史留存预测', **calc_metrics(detail, '历史留存预测')},
        {'口径': '当前v4月初回放', **calc_metrics(detail, '当前v4月初回放预测')},
    ])

    by_month = []
    for forecast_col, name in [('历史留存预测', '历史留存预测'), ('当前v4月初回放预测', '当前v4月初回放')]:
        for r in grouped_metrics(detail, forecast_col, '月份'):
            by_month.append({'口径': name, **r})
    by_month = normalise_metric_rows(by_month)

    by_bucket = []
    for forecast_col, name in [('历史留存预测', '历史留存预测'), ('当前v4月初回放预测', '当前v4月初回放')]:
        for r in grouped_metrics(detail, forecast_col, '实际销量档位'):
            by_bucket.append({'口径': name, **r})
    by_bucket = normalise_metric_rows(by_bucket)

    top_errors = sorted(
        detail,
        key=lambda r: (r['v4回放绝对误差'], r['实际销量']),
        reverse=True,
    )[:200]

    source_check = []
    for d in months:
        a = sum(v for (_spu, _shop, md), v in actual_shop.items() if md == d)
        p = sum(v for (_spu, _shop, md), v in persisted_shop.items() if md == d)
        r = sum(v for (_spu, _shop, md), v in replay_shop.items() if md == d)
        source_check.append({
            '月份': d.strftime('%Y-%m'),
            '实际销量_销量统计_msku月度': a,
            '历史留存预测': p,
            '当前v4月初回放预测': r,
        })

    out_dir = Path(args.output_dir)
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    base = out_dir / f'销量预测回测_SPU月_{stamp}'
    write_csv(base.with_name(base.name + '_明细.csv'), detail)
    write_csv(base.with_name(base.name + '_总览.csv'), overall)
    write_csv(base.with_name(base.name + '_按月份.csv'), by_month)
    write_csv(base.with_name(base.name + '_按销量档位.csv'), by_bucket)
    write_csv(base.with_name(base.name + '_TOP误差.csv'), top_errors)
    write_csv(base.with_name(base.name + '_预测方法.csv'), method_rows)
    write_csv(base.with_name(base.name + '_源数据校验.csv'), source_check)

    xlsx_path = base.with_suffix('.xlsx')
    write_xlsx(xlsx_path, [
        ('总览', overall),
        ('按月份', by_month),
        ('按销量档位', by_bucket),
        ('SPU月明细', detail),
        ('TOP误差', top_errors),
        ('预测方法', method_rows),
        ('源数据校验', source_check),
    ])

    print('\n=== 总览 ===')
    for r in overall:
        print(r)
    print('\n=== 按月份 ===')
    for r in by_month:
        print(r)
    print('\n输出目录:', out_dir.resolve())
    if Workbook is not None:
        print('Excel:', xlsx_path.resolve())
    else:
        print('openpyxl 不可用，已输出 CSV，未生成 Excel')

    print('\n重要口径说明:')
    print('1) 历史留存预测不是不可变的“月初快照”；生产任务会在目标月内反复覆盖当月预测。')
    print('2) 当前v4月初回放仅加载目标月之前的完整月销量，避免把目标月实际销量泄漏到预测输入。')
    print('3) 实际销量使用 `销量统计_msku月度`，与当前预测输入保持同源；如需用结算订单口径复核，请单独做第二真值口径。')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
