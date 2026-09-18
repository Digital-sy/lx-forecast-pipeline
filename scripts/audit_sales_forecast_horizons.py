#!/usr/bin/python3
# -*- coding: utf-8 -*-
"""Forecast audit by target month and forecast horizon (H0-H3).

This script is intentionally read-only. It replays the current production v4
forecast from historical month-start snapshots, so H1/H2/H3 exercise the same
multi-month recursive path (including the seasonal floor) used by procurement.
It also compares v4 with simple SPU-level baselines.
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from common.database import db_cursor
from jobs.feishu.forecast_sales_improved import compute_forecast_for_shop, load_spu_season_map
from jobs.feishu.generate_forecast_comparison import EXCLUDED_SHOPS, extract_spu_from_sku

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except Exception:
    Workbook = None
    Font = None


def add_months(d: date, delta: int) -> date:
    y = d.year + (d.month - 1 + delta) // 12
    m = (d.month - 1 + delta) % 12 + 1
    return date(y, m, 1)


def parse_months(text: str) -> List[date]:
    out = []
    for raw in text.split(','):
        raw = raw.strip()
        if raw:
            out.append(datetime.strptime(raw, '%Y-%m').date().replace(day=1))
    return sorted(set(out))


def sales_label(d: date) -> str:
    return f"{str(d.year)[-2:]}年{d.month}月销量"


def forecast_label(d: date) -> str:
    return f"{str(d.year)[-2:]}年{d.month}月预计销量"


def parse_stat_month(v: Any) -> Optional[date]:
    if v is None:
        return None
    if isinstance(v, datetime):
        return date(v.year, v.month, 1)
    if isinstance(v, date):
        return date(v.year, v.month, 1)
    try:
        x = datetime.strptime(str(v)[:10], '%Y-%m-%d')
        return date(x.year, x.month, 1)
    except Exception:
        return None


def table_exists(name: str) -> bool:
    with db_cursor() as cur:
        cur.execute("""SELECT COUNT(*) cnt FROM information_schema.TABLES
                       WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s""", (name,))
        return int((cur.fetchone() or {}).get('cnt', 0) or 0) > 0


def column_exists(table: str, col: str) -> bool:
    with db_cursor() as cur:
        cur.execute("""SELECT COUNT(*) cnt FROM information_schema.COLUMNS
                       WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s AND COLUMN_NAME=%s""",
                    (table, col))
        return int((cur.fetchone() or {}).get('cnt', 0) or 0) > 0


def read_sales(start: date, end: date) -> List[Dict[str, Any]]:
    table = '销量统计_msku月度'
    if not table_exists(table):
        raise RuntimeError(f'缺少表 {table}')
    has_spu = column_exists(table, 'SPU')
    spu_sql = ', SPU' if has_spu else ''
    shop_ph = ','.join(['%s'] * len(EXCLUDED_SHOPS))
    sql = f"""
        SELECT SKU, 店铺, 统计日期, 销量 {spu_sql}
        FROM `{table}`
        WHERE 统计日期 >= %s AND 统计日期 < %s
          AND 店铺 IS NOT NULL AND 店铺 NOT IN ('', '无')
          AND SKU IS NOT NULL AND SKU NOT IN ('', '无')
          AND 店铺 NOT IN ({shop_ph})
    """
    params = [start.isoformat(), end.isoformat(), *sorted(EXCLUDED_SHOPS)]
    with db_cursor() as cur:
        cur.execute(sql, params)
        return list(cur.fetchall() or [])


def row_spu(r: Dict[str, Any]) -> str:
    source = str(r.get('SPU') or '').strip()
    sku = str(r.get('SKU') or '').strip()
    return source or extract_spu_from_sku(sku)


def build_history(rows: Sequence[Dict[str, Any]], cutoff: date) -> Dict[str, Dict[str, Dict[str, Any]]]:
    out: Dict[str, Dict[str, Dict[str, Any]]] = defaultdict(lambda: defaultdict(dict))
    for r in rows:
        d = parse_stat_month(r.get('统计日期'))
        if not d or d >= cutoff:
            continue
        shop = str(r.get('店铺') or '').strip()
        sku = str(r.get('SKU') or '').strip()
        spu = row_spu(r)
        if not shop or not sku or not spu:
            continue
        x = out[shop][sku]
        x.setdefault('SPU', spu)
        label = sales_label(d)
        x[label] = int(x.get(label, 0) or 0) + int(r.get('销量') or 0)
    return {shop: dict(skus) for shop, skus in out.items()}


def actual_spu_month(rows: Sequence[Dict[str, Any]]) -> Dict[Tuple[str, date], int]:
    out: Dict[Tuple[str, date], int] = defaultdict(int)
    for r in rows:
        d = parse_stat_month(r.get('统计日期'))
        spu = row_spu(r)
        if d and spu:
            out[(spu, d)] += int(r.get('销量') or 0)
    return dict(out)


def actual_spu_shop_month(rows: Sequence[Dict[str, Any]]) -> Dict[Tuple[str, str, date], int]:
    out: Dict[Tuple[str, str, date], int] = defaultdict(int)
    for r in rows:
        d = parse_stat_month(r.get('统计日期'))
        spu = row_spu(r)
        shop = str(r.get('店铺') or '').strip()
        if d and spu and shop:
            out[(spu, shop, d)] += int(r.get('销量') or 0)
    return dict(out)


def replay(rows: Sequence[Dict[str, Any]], snapshot: date, target: date,
           season_map: Dict[str, str]) -> Tuple[Dict[Tuple[str, date], int], List[Dict[str, Any]]]:
    h = (target.year - snapshot.year) * 12 + target.month - snapshot.month
    labels = [forecast_label(add_months(snapshot, i)) for i in range(h + 1)]
    history = build_history(rows, snapshot)
    out: Dict[Tuple[str, date], int] = defaultdict(int)
    methods: List[Dict[str, Any]] = []
    target_prefix = f'{target.month}月:'

    for shop, shop_data in history.items():
        f = compute_forecast_for_shop(
            shop_data,
            labels,
            current_date=datetime(snapshot.year, snapshot.month, 1),
            spu_season_map=season_map,
        )
        tlabel = forecast_label(target)
        for sku, vals in f.items():
            spu = str(shop_data[sku].get('SPU') or extract_spu_from_sku(sku)).strip()
            if not spu:
                continue
            qty = int(vals.get(tlabel, 0) or 0)
            out[(spu, target)] += qty
            method_all = str(vals.get('预测方法') or '')
            method = ''
            for part in method_all.split('；'):
                if part.startswith(target_prefix):
                    method = part[len(target_prefix):]
                    break
            family = method.split('(', 1)[0].split('[', 1)[0] if method else 'UNKNOWN'
            methods.append({
                '快照月': snapshot.strftime('%Y-%m'),
                '目标月': target.strftime('%Y-%m'),
                'Horizon': f'H{h}',
                '店铺': shop,
                'SPU': spu,
                'SKU': sku,
                '预测销量': qty,
                '季节': season_map.get(spu, '全年'),
                '预测方法族': family,
                '命中floor': 1 if 'floor=' in method else 0,
                '预测方法': method,
            })
    return dict(out), methods


def history_value(actual: Dict[Tuple[str, date], int], spu: str, d: date) -> int:
    return int(actual.get((spu, d), 0) or 0)


def baseline_forecasts(actual: Dict[Tuple[str, date], int], spus: Iterable[str], snapshot: date,
                       target: date) -> Dict[str, Dict[Tuple[str, date], int]]:
    result = {'上月延续': {}, '近3月均值': {}, '去年同月': {}, '同比增长基线': {}}
    last = add_months(snapshot, -1)
    recent = [add_months(snapshot, -i) for i in (1, 2, 3)]
    yoy_target = add_months(target, -12)
    yoy_recent = [add_months(d, -12) for d in recent]
    for spu in spus:
        last_val = history_value(actual, spu, last)
        recent_vals = [history_value(actual, spu, d) for d in recent]
        avg3 = int(sum(recent_vals) / 3)
        yoy = history_value(actual, spu, yoy_target)
        this3 = sum(recent_vals)
        ly3 = sum(history_value(actual, spu, d) for d in yoy_recent)
        if yoy > 0 and ly3 > 0:
            gf = max(0.3, min(2.0, this3 / ly3))
            growth = int(yoy * gf)
        elif yoy > 0:
            growth = yoy
        else:
            growth = avg3
        key = (spu, target)
        result['上月延续'][key] = last_val
        result['近3月均值'][key] = avg3
        result['去年同月'][key] = yoy
        result['同比增长基线'][key] = growth
    return result


def metric(rows: Sequence[Dict[str, Any]], col: str) -> Dict[str, Any]:
    if not rows:
        return {}
    a = [float(r['实际销量'] or 0) for r in rows]
    f = [float(r[col] or 0) for r in rows]
    err = [x - y for x, y in zip(f, a)]
    ae = [abs(x) for x in err]
    ape = [abs(x-y)/y for x, y in zip(f, a) if y > 0]
    sa = sum(a)
    nz = sum(1 for y in a if y > 0)
    hit20 = sum(1 for x, y in zip(f, a) if y > 0 and abs(x-y)/y <= .2)
    hit30 = sum(1 for x, y in zip(f, a) if y > 0 and abs(x-y)/y <= .3)
    return {
        '记录数': len(rows),
        '实际销量': int(sa),
        '预测销量': int(sum(f)),
        'Bias%': None if sa == 0 else (sum(f)-sa)/sa,
        'WAPE': None if sa == 0 else sum(ae)/sa,
        'MAPE': None if not ape else sum(ape)/len(ape),
        'MAE': sum(ae)/len(ae),
        'RMSE': math.sqrt(sum(x*x for x in err)/len(err)),
        '±20%命中率': None if nz == 0 else hit20/nz,
        '±30%命中率': None if nz == 0 else hit30/nz,
    }


def norm(rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return [{k: round(v, 6) if isinstance(v, float) else v for k, v in r.items()} for r in rows]


def write_csv(path: Path, rows: Sequence[Dict[str, Any]]) -> None:
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)


def write_xlsx(path: Path, sheets: Sequence[Tuple[str, Sequence[Dict[str, Any]]]]) -> None:
    if Workbook is None:
        return
    wb = Workbook(); wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title[:31])
        if not rows:
            ws.append(['无数据']); continue
        fields = list(rows[0].keys()); ws.append(fields)
        if Font:
            for c in ws[1]: c.font = Font(bold=True)
        for r in rows: ws.append([r.get(k) for k in fields])
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            width = min(max(len(str(c.value or '')) for c in col) + 2, 55)
            ws.column_dimensions[col[0].column_letter].width = max(width, 10)
    path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--months', required=True, help='YYYY-MM,YYYY-MM,...')
    ap.add_argument('--max-horizon', type=int, default=3)
    ap.add_argument('--output-dir', default='reports_analysis/forecast_audit')
    args = ap.parse_args()
    targets = parse_months(args.months)
    min_snapshot = add_months(min(targets), -args.max_horizon)
    history_start = add_months(min_snapshot, -15)
    end = add_months(max(targets), 1)
    print('读取销量:', history_start, '~', end)
    sales_rows = read_sales(history_start, end)
    actual = actual_spu_month(sales_rows)
    season_map = load_spu_season_map()

    detail: List[Dict[str, Any]] = []
    method_rows: List[Dict[str, Any]] = []
    summary: List[Dict[str, Any]] = []

    all_spus = sorted({spu for spu, _d in actual})
    for target in targets:
        target_spus = {spu for (spu, d), qty in actual.items() if d == target and qty != 0}
        # Also retain SPUs that v4/baselines forecast even when actual is zero by using full historical universe.
        target_spus |= set(all_spus)
        for h in range(args.max_horizon + 1):
            snapshot = add_months(target, -h)
            v4, methods = replay(sales_rows, snapshot, target, season_map)
            method_rows.extend(methods)
            base = baseline_forecasts(actual, target_spus, snapshot, target)
            keys = set((spu, target) for spu in target_spus) | set(v4)
            local: List[Dict[str, Any]] = []
            for key in sorted(keys):
                spu, d = key
                a = int(actual.get(key, 0) or 0)
                row = {
                    '目标月': target.strftime('%Y-%m'),
                    '快照月': snapshot.strftime('%Y-%m'),
                    'Horizon': f'H{h}',
                    'SPU': spu,
                    '季节': season_map.get(spu, '全年'),
                    '实际销量': a,
                    'V4预测': int(v4.get(key, 0) or 0),
                    '上月延续': int(base['上月延续'].get(key, 0) or 0),
                    '近3月均值': int(base['近3月均值'].get(key, 0) or 0),
                    '去年同月': int(base['去年同月'].get(key, 0) or 0),
                    '同比增长基线': int(base['同比增长基线'].get(key, 0) or 0),
                }
                local.append(row); detail.append(row)
            for model, col in [
                ('V4', 'V4预测'), ('上月延续', '上月延续'), ('近3月均值', '近3月均值'),
                ('去年同月', '去年同月'), ('同比增长基线', '同比增长基线')]:
                summary.append({'目标月': target.strftime('%Y-%m'), 'Horizon': f'H{h}', '模型': model, **metric(local, col)})

    summary = norm(summary)

    by_horizon: List[Dict[str, Any]] = []
    for h in range(args.max_horizon + 1):
        subset = [r for r in detail if r['Horizon'] == f'H{h}']
        for model, col in [('V4','V4预测'),('上月延续','上月延续'),('近3月均值','近3月均值'),('去年同月','去年同月'),('同比增长基线','同比增长基线')]:
            by_horizon.append({'Horizon': f'H{h}', '模型': model, **metric(subset, col)})
    by_horizon = norm(by_horizon)

    by_season: List[Dict[str, Any]] = []
    for h in range(args.max_horizon + 1):
        for season in sorted({r['季节'] for r in detail}):
            subset = [r for r in detail if r['Horizon'] == f'H{h}' and r['季节'] == season]
            by_season.append({'Horizon': f'H{h}', '季节': season, '模型': 'V4', **metric(subset, 'V4预测')})
    by_season = norm(by_season)

    method_summary: List[Dict[str, Any]] = []
    groups: Dict[Tuple[str,str,str], List[Dict[str,Any]]] = defaultdict(list)
    for r in method_rows:
        groups[(r['目标月'], r['Horizon'], r['预测方法族'])].append(r)
    for (target, h, fam), rs in sorted(groups.items()):
        qty = sum(int(x['预测销量'] or 0) for x in rs)
        floors = sum(int(x['命中floor'] or 0) for x in rs)
        method_summary.append({'目标月': target, 'Horizon': h, '预测方法族': fam, 'SKU数': len(rs), '预测销量合计': qty, 'floor_SKU数': floors, 'floor命中率': floors/len(rs) if rs else 0})
    method_summary = norm(method_summary)

    floor_summary: List[Dict[str, Any]] = []
    fg: Dict[Tuple[str,str], List[Dict[str,Any]]] = defaultdict(list)
    for r in method_rows:
        fg[(r['目标月'], r['Horizon'])].append(r)
    for (target,h), rs in sorted(fg.items()):
        floors = [x for x in rs if x['命中floor']]
        floor_summary.append({'目标月': target, 'Horizon': h, 'SKU数': len(rs), 'floor_SKU数': len(floors), 'floor命中率': len(floors)/len(rs) if rs else 0, 'floor_SKU预测量': sum(int(x['预测销量'] or 0) for x in floors)})
    floor_summary = norm(floor_summary)

    top_errors = sorted(detail, key=lambda r: abs(int(r['V4预测'])-int(r['实际销量'])), reverse=True)[:300]
    for r in top_errors:
        r['V4误差'] = int(r['V4预测']) - int(r['实际销量'])
        r['V4绝对误差'] = abs(r['V4误差'])

    out = Path(args.output_dir)
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    base = out / f'销量预测Horizon回测_{stamp}'
    write_csv(base.with_name(base.name+'_Horizon总览.csv'), by_horizon)
    write_csv(base.with_name(base.name+'_月份Horizon.csv'), summary)
    write_csv(base.with_name(base.name+'_季节.csv'), by_season)
    write_csv(base.with_name(base.name+'_方法.csv'), method_summary)
    write_csv(base.with_name(base.name+'_floor.csv'), floor_summary)
    write_csv(base.with_name(base.name+'_SPU明细.csv'), detail)
    write_csv(base.with_name(base.name+'_TOP误差.csv'), top_errors)
    xlsx = base.with_suffix('.xlsx')
    write_xlsx(xlsx, [('Horizon总览',by_horizon),('月份Horizon',summary),('季节',by_season),('方法',method_summary),('floor',floor_summary),('SPU明细',detail),('TOP误差',top_errors)])

    print('\n=== Horizon总览 ===')
    for r in by_horizon: print(r)
    print('\n=== floor ===')
    for r in floor_summary: print(r)
    print('\n=== 季节 ===')
    for r in by_season: print(r)
    print('\nExcel:', xlsx.resolve())
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
